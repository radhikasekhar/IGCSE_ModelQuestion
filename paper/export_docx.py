"""Export a Paper (or its mark scheme) to an editable Word document."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt

from paper.assembler import Paper, format_duration

_TEXT_WIDTH = Cm(16)


def _setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)


def _centered(doc: Document, text: str, size: int, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)


def _add_lines(paragraph, text: str) -> None:
    lines = text.split("\n")
    for i, line in enumerate(lines):
        paragraph.add_run(line)
        if i < len(lines) - 1:
            paragraph.add_run().add_break()


def _header(doc: Document, paper: Paper, answers: bool) -> None:
    _centered(doc, ("Mark Scheme – " if answers else "") + paper.title, 16, bold=True)
    _centered(doc, paper.subtitle, 11)
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(_TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(f"Duration: {format_duration(paper.duration_minutes)}").bold = True
    p.add_run("\t")
    p.add_run(f"Total marks: {paper.total_marks}").bold = True
    if paper.instructions and not answers:
        doc.add_paragraph().add_run("INSTRUCTIONS").bold = True
        for line in paper.instructions:
            doc.add_paragraph(line, style="List Bullet")


def _footer(doc: Document, paper: Paper) -> None:
    fp = doc.sections[0].footer.paragraphs[0]
    fp.text = f"{paper.title} · Generated {paper.created_at} · Seed {paper.seed}"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        r.font.size = Pt(8)


def build(paper: Paper, answers: bool = False) -> bytes:
    doc = Document()
    _setup(doc)
    _header(doc, paper, answers)
    _footer(doc, paper)

    for sec in paper.sections:
        if sec.title:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            run = h.add_run(sec.title)
            run.bold, run.font.size = True, Pt(13)
        for q in sec.questions:
            if q.stem_context:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.first_line_indent = Cm(-1)
                p.add_run(f"{q.number}\t").bold = True
                _add_lines(p, q.stem_context)
                if not answers:
                    for img in q.image_paths:
                        if Path(img).exists():
                            doc.add_picture(img, width=Cm(12))
                            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.first_line_indent = Cm(-1)
                p.add_run(f"{q.number}\t").bold = True
            p.paragraph_format.tab_stops.add_tab_stop(_TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
            _add_lines(p, q.text)
            p.add_run(f"\t[{q.marks}]").bold = True
            if paper.show_sources:
                ref = p.add_run(f"\n{q.source_ref}")
                ref.italic, ref.font.size = True, Pt(8)
            if answers:
                a = doc.add_paragraph()
                a.paragraph_format.left_indent = Cm(1)
                a.add_run("Answer: ").bold = True
                _add_lines(a, q.ms_answer or "Mark scheme not available in source documents.")
            elif not q.stem_context:
                for img in q.image_paths:
                    if Path(img).exists():
                        doc.add_picture(img, width=Cm(12))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    if paper.inserts and not answers:
        doc.add_page_break()
        h = doc.add_paragraph()
        h.add_run("INSERT – Source material").bold = True
        for text in paper.inserts.values():
            _add_lines(doc.add_paragraph(), text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
