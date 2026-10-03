"""Load supported files into a list of pages (text, or a rendered image for scanned pages)."""
from __future__ import annotations

import base64
import io
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf as fitz
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from PIL import Image

import config


@dataclass
class Page:
    number: int  # 1-based
    text: str = ""
    image_b64: str | None = None  # set when the page must be read with vision
    image_media_type: str = "image/png"


@dataclass
class LoadedDocument:
    path: Path
    kind: str  # pdf | docx | txt | html | image
    pages: list[Page] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(p.text) for p in self.pages)


class UnsupportedFileError(Exception):
    pass


def load(path: Path) -> LoadedDocument:
    ext = path.suffix.lower()
    if ext not in config.SUPPORTED_EXTS:
        raise UnsupportedFileError(f"Unsupported file type: {ext}")
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext in (".html", ".htm"):
        return _load_html(path)
    if ext == ".txt":
        return LoadedDocument(path, "txt", [Page(1, path.read_text(encoding="utf-8", errors="replace"))])
    return _load_image(path)


def _load_pdf(path: Path) -> LoadedDocument:
    doc = LoadedDocument(path, "pdf")
    with fitz.open(path) as pdf:
        for i, page in enumerate(pdf, start=1):
            text = page.get_text("text", sort=True)
            p = Page(i, text)
            if len(text.strip()) < config.SCANNED_PAGE_MIN_CHARS:
                # Scanned or image-only page: send a rendering to the vision model instead.
                pix = page.get_pixmap(dpi=150)
                p.image_b64 = base64.standard_b64encode(pix.tobytes("png")).decode("ascii")
            doc.pages.append(p)
    return doc


def _load_docx(path: Path) -> LoadedDocument:
    d = DocxDocument(str(path))
    lines: list[str] = [para.text for para in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    # DOCX has no fixed pages; group ~60 lines per logical page so citations stay useful.
    pages = [Page(i + 1, "\n".join(lines[j : j + 60])) for i, j in enumerate(range(0, max(len(lines), 1), 60))]
    return LoadedDocument(path, "docx", pages)


def _load_html(path: Path) -> LoadedDocument:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    return LoadedDocument(path, "html", [Page(1, soup.get_text("\n"))])


def _load_image(path: Path) -> LoadedDocument:
    # Downscale large photos/scans: the API caps images at 5 MB and gains nothing past ~2000 px.
    with Image.open(path) as img:
        img = img.convert("RGB")
        img.thumbnail((2000, 2000))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=88)
    data = base64.standard_b64encode(buf.getvalue()).decode("ascii")
    return LoadedDocument(path, "image", [Page(1, "", data, "image/jpeg")])
